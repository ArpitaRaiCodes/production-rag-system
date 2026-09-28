"""Turn an uploaded file into clean, overlapping text chunks."""

from __future__ import annotations

import csv
import io
import re
from dataclasses import dataclass
from pathlib import Path

from app.config import settings

SUPPORTED_EXTENSIONS = {".pdf", ".txt", ".md", ".markdown", ".docx", ".csv", ".json"}


class UnsupportedFileType(Exception):
    pass


class EmptyDocument(Exception):
    pass


@dataclass
class Chunk:
    text: str
    index: int
    page: int | None = None


# ---------------------------------------------------------------------------
# Extraction
# ---------------------------------------------------------------------------
def _extract_pdf(path: Path) -> list[tuple[str, int]]:
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    pages: list[tuple[str, int]] = []
    for number, page in enumerate(reader.pages, start=1):
        try:
            text = page.extract_text() or ""
        except Exception:  # a damaged page should not kill the whole upload
            text = ""
        if text.strip():
            pages.append((text, number))
    return pages


def _extract_docx(path: Path) -> list[tuple[str, int]]:
    import docx  # python-docx

    document = docx.Document(str(path))
    parts = [p.text for p in document.paragraphs if p.text.strip()]
    for table in document.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    return [("\n".join(parts), None)] if parts else []


def _extract_csv(path: Path) -> list[tuple[str, int]]:
    raw = path.read_text(encoding="utf-8", errors="ignore")
    reader = csv.reader(io.StringIO(raw))
    rows = list(reader)
    if not rows:
        return []
    header = rows[0]
    lines = []
    for row in rows[1:]:
        pairs = [f"{h}: {v}" for h, v in zip(header, row) if v.strip()]
        if pairs:
            lines.append("; ".join(pairs))
    return [("\n".join(lines), None)] if lines else []


def _extract_plain(path: Path) -> list[tuple[str, int]]:
    text = path.read_text(encoding="utf-8", errors="ignore")
    return [(text, None)] if text.strip() else []


def extract_text(path: Path) -> list[tuple[str, int | None]]:
    suffix = path.suffix.lower()
    if suffix not in SUPPORTED_EXTENSIONS:
        raise UnsupportedFileType(
            f"{suffix or 'this file'} is not supported. "
            f"Supported types: {', '.join(sorted(SUPPORTED_EXTENSIONS))}"
        )
    if suffix == ".pdf":
        return _extract_pdf(path)
    if suffix == ".docx":
        return _extract_docx(path)
    if suffix == ".csv":
        return _extract_csv(path)
    return _extract_plain(path)


# ---------------------------------------------------------------------------
# Cleaning + chunking
# ---------------------------------------------------------------------------
def clean(text: str) -> str:
    text = text.replace("\x00", " ").replace("\u00ad", "")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def _split_sentences(text: str) -> list[str]:
    parts = re.split(r"(?<=[.!?])\s+(?=[A-Z0-9\"'(])|\n{2,}", text)
    return [p.strip() for p in parts if p and p.strip()]


def chunk_text(
    text: str,
    chunk_size: int | None = None,
    overlap: int | None = None,
) -> list[str]:
    """Greedy sentence packing with a character budget and a character overlap.

    Sentence-aware packing keeps chunks readable, which matters because the
    retrieved chunk is what the model actually reads.
    """
    chunk_size = chunk_size or settings.chunk_size
    overlap = overlap if overlap is not None else settings.chunk_overlap
    overlap = min(overlap, max(chunk_size - 50, 0))

    text = clean(text)
    if not text:
        return []

    sentences = _split_sentences(text)
    chunks: list[str] = []
    current = ""

    for sentence in sentences:
        # A single very long sentence gets hard-split.
        while len(sentence) > chunk_size:
            if current:
                chunks.append(current.strip())
                current = ""
            chunks.append(sentence[:chunk_size].strip())
            sentence = sentence[chunk_size - overlap :]

        if len(current) + len(sentence) + 1 <= chunk_size:
            current = f"{current} {sentence}".strip()
        else:
            if current:
                chunks.append(current.strip())
            tail = current[-overlap:] if overlap and current else ""
            current = f"{tail} {sentence}".strip()

    if current.strip():
        chunks.append(current.strip())

    return [c for c in chunks if len(c) > 30]


def build_chunks(path: Path) -> list[Chunk]:
    """Read a file and return the chunks that will be embedded."""
    pages = extract_text(path)
    if not pages:
        raise EmptyDocument(
            "No text could be read from this file. If it is a scanned PDF, "
            "run OCR on it first (see GUIDE.md)."
        )

    chunks: list[Chunk] = []
    index = 0
    for text, page_number in pages:
        for piece in chunk_text(text):
            chunks.append(Chunk(text=piece, index=index, page=page_number))
            index += 1

    if not chunks:
        raise EmptyDocument("The file was read but contained no usable text.")
    return chunks
